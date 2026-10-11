"""
excel_data_provider.py — Local Excel fallback for inventory, FAQ, and order data.

When Google Sheets credentials are NOT configured, the application reads
menu and FAQ data from the local Food_Delivery_System.xlsx file and
writes orders to a local JSON file.

This lets the WhatsApp bot work end-to-end in development without needing
a Google Cloud service account or live Sheets.
"""
from __future__ import annotations

import json
import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

import openpyxl

logger = logging.getLogger(__name__)

# Resolve the data directory relative to the project root
_PROJECT_ROOT = Path(__file__).resolve().parents[3]
_EXCEL_PATH = _PROJECT_ROOT / "data" / "Food_Delivery_System.xlsx"
_ORDERS_JSON_PATH = _PROJECT_ROOT / "data" / "local_orders.json"


def _excel_exists() -> bool:
    return _EXCEL_PATH.exists()


def is_sheets_configured() -> bool:
    """Return True if Google Sheets credentials and IDs are properly configured."""
    from app.config import settings

    # Check for service account file
    sa_path = settings.google_service_account_json_path
    if not sa_path:
        return False

    project_root = Path(__file__).resolve().parents[3]
    creds_path = Path(sa_path)
    if not creds_path.is_absolute():
        creds_path = project_root / creds_path
    if not creds_path.exists():
        return False

    # Check for real spreadsheet IDs (not placeholders)
    inv_id = settings.inventory_spreadsheet_id or ""
    if not inv_id or "your_" in inv_id.lower():
        return False

    return True


class ExcelDataProvider:
    """
    Reads data from the local Food_Delivery_System.xlsx file.

    Thread-safety: openpyxl workbooks are read-only here and we reload
    on each call, so this is safe for concurrent access.
    """

    def __init__(self, excel_path: Path = _EXCEL_PATH) -> None:
        self._path = excel_path

    def read_inventory(self) -> list[dict[str, Any]]:
        """Read the Inventory sheet as a list of dicts (same format as gspread)."""
        if not self._path.exists():
            logger.warning("Excel file not found at %s", self._path)
            return []

        wb = openpyxl.load_workbook(str(self._path), read_only=True, data_only=True)
        try:
            ws = wb["Inventory"]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return []

            headers = [str(h).strip().lower() if h else "" for h in rows[0]]
            result = []
            for row in rows[1:]:
                record = {}
                for i, val in enumerate(row):
                    if i < len(headers) and headers[i]:
                        record[headers[i]] = val if val is not None else ""
                if record.get("item_id"):
                    result.append(record)

            logger.info("Loaded %d inventory items from local Excel.", len(result))
            return result
        finally:
            wb.close()

    def read_faqs_raw_column(self) -> list[str]:
        """Read the faqs sheet column A as a list of strings (single-column format)."""
        if not self._path.exists():
            logger.warning("Excel file not found at %s", self._path)
            return []

        wb = openpyxl.load_workbook(str(self._path), read_only=True, data_only=True)
        try:
            ws = wb["faqs"]
            values = []
            for row in ws.iter_rows(min_col=1, max_col=1, values_only=True):
                val = row[0]
                values.append(str(val) if val is not None else "")
            logger.info("Loaded %d FAQ rows from local Excel.", len(values))
            return values
        finally:
            wb.close()


class LocalOrderStore:
    """
    Writes orders to a local JSON file when Google Sheets is not available.
    Each order is appended as a line in a JSON-lines file.
    """

    def __init__(self, path: Path = _ORDERS_JSON_PATH) -> None:
        self._path = path

    def append_order(self, row: list[Any]) -> None:
        """Append an order row to the local JSON file."""
        order_record = {
            "order_id": str(row[0]) if len(row) > 0 else "",
            "customer_name": str(row[1]) if len(row) > 1 else "",
            "phone": str(row[2]) if len(row) > 2 else "",
            "item_name": str(row[3]) if len(row) > 3 else "",
            "quantity": str(row[4]) if len(row) > 4 else "",
            "unit_price": str(row[5]) if len(row) > 5 else "",
            "total_amount": float(row[6]) if len(row) > 6 else 0,
            "delivery_address": str(row[7]) if len(row) > 7 else "",
            "payment_method": str(row[8]) if len(row) > 8 else "",
            "order_status": str(row[9]) if len(row) > 9 else "",
            "order_date": str(row[10]) if len(row) > 10 else "",
            "saved_at": datetime.now(timezone.utc).isoformat(),
        }

        # Append as JSON-lines
        with open(self._path, "a", encoding="utf-8") as f:
            f.write(json.dumps(order_record, ensure_ascii=False) + "\n")

        logger.info("Order %s saved to local file.", order_record["order_id"])

    def find_order_by_id(self, order_id: str) -> Optional[dict]:
        """Check if an order ID already exists in the local store."""
        if not self._path.exists():
            return None
        with open(self._path, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                try:
                    record = json.loads(line)
                    if record.get("order_id") == order_id:
                        return record
                except json.JSONDecodeError:
                    continue
        return None


# Module-level singletons
excel_provider = ExcelDataProvider()
local_order_store = LocalOrderStore()
